import importlib.util
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from PIL import Image
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from privacy_core import (Detector, Finding, PrivacyError, Registry, draw_redactions,
                          key_bytes, read_vault, resolve_findings, words_from_tsv)
from anonymize_document import approve, prepare, restore, verified_vault


def empty_ner(text):
    return SimpleNamespace(ents=[])


class PrivacyDetectionTests(unittest.TestCase):
    def test_rules_fields_known_values_and_local_ner_combine(self):
        def nlp(text):
            start = text.find("Alice Example")
            return SimpleNamespace(ents=[] if start < 0 else [SimpleNamespace(
                start_char=start, end_char=start + len("Alice Example"), label_="PERSON")])
        detector = Detector(known_values=["Hidden Street"], nlp=nlp)
        text = "Alice Example\nEmail: alice@example.test\nPassport number: AB123456\nHidden Street"
        findings = detector.find(text)
        registry = Registry("AABBCCDD")
        redacted, _ = registry.redact(text, findings)
        for value in ("Alice Example", "alice@example.test", "AB123456", "Hidden Street"):
            self.assertNotIn(value, redacted)
        self.assertEqual(registry.restore(redacted), text)

    def test_partial_overlaps_are_unioned(self):
        self.assertEqual(resolve_findings([Finding(0, 8, "PERSON", "ner"), Finding(5, 12, "FIELD", "field")]),
                         [Finding(0, 12, "PRIVATE", "combined")])

    def test_existing_tokens_not_redetected(self):
        self.assertEqual(Detector(nlp=empty_ner).find("ICATPII_AABBCCDD_PERSON_0001"), [])

    def test_unvalidated_language_fails_closed(self):
        with self.assertRaises(PrivacyError):
            Detector(language="ar", nlp=empty_ner)

    def test_latin_language_field_profiles(self):
        for language, label in [("de", "Vorname"), ("fr", "Prénom"), ("es", "Nombre"),
                                ("it", "Cognome"), ("pt", "Sobrenome")]:
            with self.subTest(language=language):
                text = label + ": Alice Example"
                redacted, _ = Registry().redact(text, Detector(language=language, nlp=empty_ner).find(text))
                self.assertNotIn("Alice Example", redacted)

    def test_repeated_values_use_same_token_but_documents_have_distinct_namespaces(self):
        registry = Registry()
        token = registry.register("Alice", "PERSON")
        self.assertEqual(token, registry.register("Alice", "KNOWN"))
        self.assertNotEqual(token, Registry().register("Alice", "PERSON"))
        with self.assertRaises(PrivacyError):
            Registry().restore(token)

    def test_ocr_coordinates_and_multiline_offsets(self):
        tsv = ("level\tpage_num\tblock_num\tpar_num\tline_num\tword_num\tleft\ttop\twidth\theight\tconf\ttext\n"
               "5\t1\t1\t1\t1\t1\t10\t10\t30\t20\t98\tAlice\n"
               "5\t1\t1\t1\t1\t2\t50\t10\t60\t20\t95\tExample\n"
               "5\t1\t1\t1\t2\t1\t10\t50\t60\t20\t65\t12345678\n")
        text, words = words_from_tsv(tsv)
        self.assertEqual(text, "Alice Example\n12345678")
        self.assertEqual(text[words[2]["start"]:words[2]["end"]], "12345678")
        self.assertEqual(words[2]["box"], [10, 50, 70, 70])

    def test_image_pixels_overwritten_and_metadata_removed(self):
        image = Image.new("RGB", (500, 100), "red")
        image.info["secret"] = "Alice"
        words = [{"start": 0, "end": 5, "box": [20, 20, 250, 50], "line": [1, 1, 1, 1]}]
        clean, masks = draw_redactions(image, words, [(Finding(0, 5, "PERSON", "ner"), "ICATPII_AABBCCDD_PERSON_0001")])
        self.assertEqual(clean.getpixel((245, 48)), (255, 255, 255))
        self.assertEqual(clean.getpixel((499, 99)), (255, 0, 0))
        self.assertEqual(clean.info, {})
        self.assertTrue(masks)
        with self.assertRaises(PrivacyError):
            draw_redactions(image, [], [], [{"box": [-1, 0, 10, 10], "token": "x"}])


@unittest.skipUnless(importlib.util.find_spec("cryptography"), "Optional privacy dependencies are not installed")
class PrivacyWorkflowTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.root = Path(self.temporary.name)
        self.key = key_bytes(self.root / "keys" / "privacy.key", create=True)
        self.input = self.root / "source.png"
        Image.new("RGB", (500, 150), "white").save(self.input)
        self.directory = self.root / "prepared"
        self.ocr_text = "Name: Alice Example\nPassport: AB123456"
        self.words = [
            {"start": 0, "end": 5, "box": [10, 20, 70, 40], "line": [1, 1, 1, 1], "confidence": 99},
            {"start": 6, "end": 19, "box": [90, 20, 300, 40], "line": [1, 1, 1, 1], "confidence": 99},
            {"start": 20, "end": 29, "box": [10, 60, 80, 80], "line": [1, 1, 1, 2], "confidence": 99},
            {"start": 30, "end": 38, "box": [100, 60, 300, 80], "line": [1, 1, 1, 2], "confidence": 50},
        ]

    def preparation(self, **options):
        # Synthetic inputs only; real OCR is verified in a separate CLI smoke check.
        with patch("anonymize_document.local_ocr", return_value=(self.ocr_text, self.words)):
            return prepare(self.input, self.directory, self.key, detector=Detector(nlp=empty_ner), **options)

    def test_approval_requires_review_and_pdf_has_no_original_text_layer(self):
        self.preparation()
        with self.assertRaises(PrivacyError):
            approve(self.directory, self.key)
        pdf = approve(self.directory, self.key, reviewed=True)
        import pymupdf
        with pymupdf.open(pdf) as document:
            self.assertEqual(len(document), 1)
            self.assertEqual(document[0].get_text(), "")
            self.assertFalse(document.metadata.get("title"))
        self.assertNotIn(b"Alice Example", (self.directory / "vault.enc").read_bytes())
        self.assertNotIn("Alice Example", (self.directory / "page-001.txt").read_text())
        if os.name != "nt":
            self.assertEqual((self.directory / "vault.enc").stat().st_mode & 0o777, 0o600)

    def test_modified_page_or_review_cannot_be_approved(self):
        self.preparation()
        (self.directory / "page-001.png").write_bytes(b"changed")
        with self.assertRaises(PrivacyError):
            approve(self.directory, self.key, reviewed=True)
        self.assertFalse((self.directory / "approved.pdf").exists())

    def test_missing_label_blocks_approval_and_wider_region_can_fix_it(self):
        self.ocr_text = "Name: Al"
        self.words = [{"start": 0, "end": 5, "box": [10, 20, 70, 40], "line": [1, 1, 1, 1], "confidence": 99},
                      {"start": 6, "end": 8, "box": [90, 20, 110, 40], "line": [1, 1, 1, 1], "confidence": 99}]
        self.preparation()
        with self.assertRaises(PrivacyError):
            approve(self.directory, self.key, reviewed=True)
        self.directory = self.root / "prepared-wider"
        self.preparation(regions=[{"page": 0, "box": [85, 15, 350, 50]}])
        self.assertTrue(approve(self.directory, self.key, reviewed=True).exists())

    def test_cannot_reuse_output_or_overwrite_key(self):
        self.preparation()
        with self.assertRaises(PrivacyError):
            self.preparation()
        with self.assertRaises(PrivacyError):
            key_bytes(self.root / "keys" / "privacy.key", create=True)

    def test_manual_region_cleans_text_sidecar_as_well_as_image(self):
        self.ocr_text = "ObscureSensitiveValue"
        self.words = [{"start": 0, "end": len(self.ocr_text), "box": [20, 20, 300, 50],
                       "line": [1, 1, 1, 1], "confidence": 50}]
        self.preparation(regions=[{"page": 0, "box": [10, 10, 310, 60]}])
        self.assertNotIn(self.ocr_text, (self.directory / "page-001.txt").read_text())
        self.assertEqual(len(verified_vault(self.directory, self.key)["entries"]), 1)

    def test_failed_ocr_blocks_approval_without_exposing_ocr_error_contents(self):
        with patch("anonymize_document.local_ocr", return_value=("", [])):
            with self.assertRaises(PrivacyError):
                prepare(self.input, self.directory, self.key, detector=Detector(nlp=empty_ner))
        self.assertTrue((self.directory / "FAILED").exists())
        self.assertFalse((self.directory / "approved.pdf").exists())

    def test_restore_docx_across_runs_with_local_replacement(self):
        self.preparation()
        approve(self.directory, self.key, reviewed=True)
        vault = verified_vault(self.directory, self.key)
        tokens = list(vault["entries"])
        document = Document()
        for token in tokens:
            paragraph = document.add_paragraph("Before ")
            paragraph.add_run(token[:12]).bold = True
            paragraph.add_run(token[12:])
            paragraph.add_run(" after")
        document.add_paragraph(tokens[0] + " / " + tokens[0])
        source = self.root / "translated.docx"
        document.save(source)
        destination = self.root / "restored.docx"
        report = restore(source, destination, self.directory, self.key, {tokens[0]: "Approved spelling"})
        self.assertFalse(report["missingTokens"])
        text = "\n".join(p.text for p in Document(destination).paragraphs)
        self.assertIn("Before Approved spelling after", text)
        self.assertIn("Approved spelling / Approved spelling", text)
        self.assertNotIn("ICATPII_", text)

    def test_missing_unknown_or_tampered_placeholders_block_restoration(self):
        self.preparation()
        approve(self.directory, self.key, reviewed=True)
        tokens = list(verified_vault(self.directory, self.key)["entries"])
        source = self.root / "translated.txt"
        source.write_text(tokens[0])
        with self.assertRaises(PrivacyError):
            restore(source, self.root / "restored.txt", self.directory, self.key)
        self.assertFalse((self.root / "restored.txt").exists())
        source.write_text("ICATPII_DEADBEEF_PERSON_0001")
        with self.assertRaises(PrivacyError):
            restore(source, self.root / "restored.txt", self.directory, self.key)
        source.write_text(" ".join(tokens))
        (self.directory / "approved.pdf").write_bytes(b"tampered")
        with self.assertRaises(PrivacyError):
            restore(source, self.root / "restored.txt", self.directory, self.key)


if __name__ == "__main__":
    unittest.main()
