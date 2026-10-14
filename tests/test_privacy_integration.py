"""Opt-in real local OCR/NER acceptance check; uses only a synthetic document."""
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pymupdf
from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from anonymize_document import approve, prepare, restore, verified_vault
from privacy_core import Detector, key_bytes


@unittest.skipUnless(os.environ.get("ICAT_PRIVACY_INTEGRATION_TEST") == "1", "Run npm run verify:privacy for real OCR/NER")
class LocalPrivacyAcceptanceTests(unittest.TestCase):
    def test_real_ocr_redaction_and_docx_restoration_without_python_network(self):
        detector = Detector("en")  # Requires the installed model; never downloads one.
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            source = root / "source.pdf"
            document = pymupdf.open()
            page = document.new_page()
            page.insert_text((45, 55), "SYNTHETIC PRIVACY TEST", fontsize=16)
            page.insert_text((45, 95), "Name: Alice Example\nEmail: alice@example.test\nPassport number: AB123456\n"
                             "Address: 42 Hidden Street\nPhone: +1 202 555 0123\nDate of birth: 01/02/1990", fontsize=14)
            document.set_metadata({"author": "Private Example Author"})
            document.save(source)
            document.close()
            key = key_bytes(root / "keys" / "local.key", create=True)
            directory = root / "prepared"
            with patch("socket.socket.connect", side_effect=AssertionError("Network access is forbidden")), \
                 patch("socket.create_connection", side_effect=AssertionError("Network access is forbidden")):
                prepare(source, directory, key, detector=detector)
                text = (directory / "page-001.txt").read_text()
                for value in ["Alice Example", "alice@example.test", "AB123456", "Hidden Street", "555 0123", "01/02/1990"]:
                    self.assertNotIn(value, text)
                pdf_path = approve(directory, key, reviewed=True)
                with pymupdf.open(pdf_path) as clean:
                    self.assertEqual(clean[0].get_text(), "")
                    self.assertFalse(clean.metadata.get("author"))
                vault = verified_vault(directory, key)
                translated = Document()
                for token in vault["entries"]:
                    paragraph = translated.add_paragraph("Translated field: ")
                    paragraph.add_run(token[:15])
                    paragraph.add_run(token[15:])
                translated.save(root / "translation.docx")
                report = restore(root / "translation.docx", root / "restored.docx", directory, key)
                self.assertFalse(report["missingTokens"])
                result = "\n".join(p.text for p in Document(root / "restored.docx").paragraphs)
                self.assertIn("Alice Example", result)
                self.assertNotIn("ICATPII_", result)


if __name__ == "__main__":
    unittest.main()
